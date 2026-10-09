import openpyxl
filepath = 'D:/Kyaw Zay Ya Written Project/animal/frontend/src/data/WDN_Survey_Report_2026-10-09 (3).xlsx'
wb = openpyxl.load_workbook(filepath)
print("Sheet names:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\nSheet: {name}")
    print(f"Rows: {ws.max_row}, Columns: {ws.max_column}")
    headers = [cell.value for cell in ws[1]]
    print(f"Total headers: {len(headers)}")
    print(f"First 50 headers: {headers[:50]}")
    print(f"All headers count: {len(headers)}")
    for i, h in enumerate(headers, 1):
        print(f"  {i}. {h}")